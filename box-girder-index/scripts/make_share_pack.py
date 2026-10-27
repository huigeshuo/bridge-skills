# -*- coding: utf-8 -*-
"""生成 box-girder-index 的「对外分享包」（脱敏版）。

用法：
  python make_share_pack.py --outdir "<目标目录>"

做的事：
  1) 复制 scripts/extract_indicators.py、SKILL.md、references/指标口径与校核规则.md
  2) 把其中的真实工程名/互通名/桥名替换为代号，并清空 NAME_ALIAS
  3) 把本机绝对路径（含用户名）替换为占位符
  4) 生成一份可直接被脚本解析的「输入表模板.xlsx」
  5) 复制对外精简版 HTML 作为成品示例
"""
import argparse
import io
import os
import re
import shutil

SKILL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# 真实名 -> 代号（长词在前）
ALIAS = [
    ("武大高速", "WD高速"),
    ("武大", "WD"),
    ("汉口北互通钢束含量汇总表-szh.xlsx", "HKB互通钢束含量汇总表.xlsx"),
    ("汉口北互通", "HKB互通"),
    ("汉口北大道高架", "HKBDDGJ"),
    ("汉口北", "HKB"),
    ("滠口高架", "SKGJ"),
    ("喻家咀", "YJJ"),
    ("小菜湖", "XCH"),
    ("肖家吉", "XJJ"),
    ("刘集", "LJ"),
    ("青峰", "QF"),
    ("涂店", "TD"),
    ("主线右幅桥", "ZXYFQ"),
    ("主线左幅桥", "ZXZFQ"),
    ("A匝道桥", "AZDQ"),
]

# 路径 / 用户名脱敏
PATH_FIX = [
    ("C:/Users/Administrator/.workbuddy/binaries/python/envs/default/Scripts/python.exe", "<python>"),
    ("C:/Users/Administrator/.workbuddy/binaries/python/envs/default/Scripts/python", "<python>"),
    ("C:\\Users\\Administrator\\.workbuddy\\binaries\\python\\envs\\default", "<你的 Python 环境>"),
    ("C:/Users/Administrator/.workbuddy/skills/box-girder-index", "<skill>"),
    ("F:/SZH 文件夹/228国道斜拉桥/02 参考图/02 武大技术指标/", ""),
    ("F:\\SZH 文件夹\\武大箱梁技术指标.xlsx", "你的汇总表.xlsx"),
    ("F:/SZH 文件夹/Obsidian Vault bridge/13 工具开发/", "<你的知识库>/13 工具开发/"),
]

# NAME_ALIAS 整块替换（脚本里）
OLD_ALIAS = '''NAME_ALIAS = {
    "武大高速": "WD高速",   # 长词在前，避免被短词先吃掉
    "武大":     "WD",
    "汉口北":   "HKB",
}'''
NEW_ALIAS = '''# 默认留空：本分享包不含任何真实工程名。
# 如果你自己的表名/互通名也需要脱敏，在这里登记，长词必须排在短词前。
# 例：{"某某高速": "MM高速", "某某": "MM"}
NAME_ALIAS = {}'''


def scrub(text: str) -> str:
    # 先清路径（含真实工程名），再清工程名；最后正则兜底，防止顺序问题漏网
    for a, b in PATH_FIX:
        text = text.replace(a, b)
    for a, b in ALIAS:
        text = text.replace(a, b)
    text = re.sub(r'[A-Za-z]:[\\/](?:[^\\\s"\']*[\\/])*[^\\\s"\']*\.xlsx?',
                  '你的汇总表.xlsx', text)
    text = re.sub(r'[A-Za-z]:[\\/]SZH[^\s"\']*', '<你的目录>/', text)
    return text


def build_script() -> str:
    p = os.path.join(SKILL, "scripts", "extract_indicators.py")
    s = io.open(p, encoding="utf-8").read()
    assert OLD_ALIAS in s, "NAME_ALIAS 块未匹配，脚本结构可能已变"
    s = s.replace(OLD_ALIAS, NEW_ALIAS)
    return scrub(s)


def build_skill_md() -> str:
    p = os.path.join(SKILL, "SKILL.md")
    s = io.open(p, encoding="utf-8").read()
    s = scrub(s)
    # 第六节案例标题改口径，避免读者以为还带真名
    s = s.replace("## 六、真实案例（2026-10-22，WD高速 HKB互通）",
                  "## 六、实测案例（已代号化，仅用于说明口径与校核效果）")
    # 去掉「第七节代号对照表」相关提示（本包不含对照表）
    s = s.replace("**该对照表是找回原名的唯一钥匙，别删。**",
                  "（本分享包已统一脱敏，不含原名对照表。）")
    return s


def build_reference() -> str:
    p = os.path.join(SKILL, "references", "指标口径与校核规则.md")
    s = io.open(p, encoding="utf-8").read()
    return scrub(s)


SOURCE_ALIAS = [
    ("汉口北互通钢束含量汇总表-szh.xlsx", "HKB互通钢束含量汇总表.xlsx"),
    ("武大箱梁技术指标.xlsx", "WD箱梁技术指标.xlsx"),
]


def build_sample_lib(json_path: str):
    """生成可外发的样本库：剥离原始中文字段 + 源文件名代号化，保留全部数值指标。

    与网页看板同源，可被 extract_indicators.py 作为 --lib 直接读取，也能喂给 AI 做对话式检索。
    """
    import json

    src = os.path.join(SKILL, "data", "box_girder_indicators.json")
    records = json.load(io.open(src, encoding="utf-8"))
    safe = []
    for r in records:
        d = {k: v for k, v in r.items() if k not in ("route_src", "bridge_src")}
        sf = d.get("source_file") or ""
        for a, b in SOURCE_ALIAS:
            sf = sf.replace(a, b)
        d["source_file"] = scrub(sf)
        # 兜底：任何字符串字段再过一遍脱敏
        for k, v in list(d.items()):
            if isinstance(v, str):
                d[k] = scrub(v)
        safe.append(d)
    json.dump(safe, io.open(json_path, "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    return len(safe)


def build_template(xlsx_path: str):
    """生成一份脚本能直接解析的输入表模板（含区块标题行 + 表头 + 2 行示例）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "现浇箱梁"

    # 第 1 行：区块标题行（横向合并 >= 4 列才会被识别为区块标题）
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    c = ws.cell(row=1, column=1, value="某某互通(现浇)")
    c.font = Font(bold=True, size=12)
    c.alignment = Alignment(horizontal="center")

    header = ["桥名", "互通", "联号", "跨径", "联长(m)",
              "起点宽度(m)", "终点宽度(m)", "桥梁面积(m2)",
              "砼数量(m3)", "钢绞线数量(kg)", "备注"]
    ws.append(header)

    rows = [
        ["A匝道桥", "某某互通", 1, "3*30", 90.0, 9.0, 9.0, 810.0, 600.0, 16000.0, "示例行"],
        ["A匝道桥", "某某互通", 2, "25+35+25", 85.0, 9.0, 12.5, 911.0, 700.0, 19000.0, "示例行·变宽"],
    ]
    for r in rows:
        ws.append(r)

    for i, w in enumerate([12, 12, 8, 16, 10, 12, 12, 14, 12, 16, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(xlsx_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--outdir", required=True)
    a = ap.parse_args()
    out = a.outdir
    os.makedirs(os.path.join(out, "scripts"), exist_ok=True)
    os.makedirs(os.path.join(out, "references"), exist_ok=True)

    io.open(os.path.join(out, "scripts", "extract_indicators.py"), "w",
            encoding="utf-8").write(build_script())
    io.open(os.path.join(out, "SKILL.md"), "w", encoding="utf-8").write(build_skill_md())
    io.open(os.path.join(out, "references", "指标口径与校核规则.md"), "w",
            encoding="utf-8").write(build_reference())
    build_template(os.path.join(out, "输入表模板.xlsx"))
    os.makedirs(os.path.join(out, "data"), exist_ok=True)
    n = build_sample_lib(os.path.join(out, "data", "指标库-样本数据.json"))
    print("[样本库] 已生成 %d 条（已剥离原名与真实源文件名）" % n)

    src_html = os.path.join(SKILL, "data", "out", "指标库-对外精简版.html")
    if os.path.exists(src_html):
        shutil.copy(src_html, os.path.join(out, "箱梁技术指标速查-辉歌说.html"))

    print("[完成] 对外分享包 ->", out)
    for f in sorted(os.listdir(out)):
        p = os.path.join(out, f)
        if os.path.isdir(p):
            for g in sorted(os.listdir(p)):
                print("   ", os.path.join(f, g))
        else:
            print("   ", f)


if __name__ == "__main__":
    main()
